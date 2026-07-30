"""Build a single Excel annotation workbook from an annotation packet."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter, quote_sheetname
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # pragma: no cover - exercised only without the extra
    raise SystemExit(
        "Install the workbook dependency first: "
        "python -m pip install -r "
        "benchmarks/contexttrace_unseen_v1/requirements-annotation-xlsx.txt"
    ) from exc


CLAIM_SLOTS_PER_CASE = 20
MAX_EXCEL_TEXT = 32_000

OPTIONS: dict[str, list[str]] = {
    "propositional": ["yes", "no"],
    "claim_verdict": [
        "supported",
        "partially_supported",
        "unsupported",
        "contradicted",
        "unverifiable",
    ],
    "failure_label": [
        "none",
        "retrieval_miss",
        "context_selection_error",
        "citation_mismatch",
        "answer_overreach",
        "contradiction",
        "insufficient_evidence",
        "should_have_abstained",
        "source_condition_failure",
    ],
    "primary_root_cause": [
        "none",
        "retrieval_miss",
        "reranking_failure",
        "chunking_issue",
        "corpus_gap",
        "stale_or_superseded_source",
        "noncanonical_or_low_authority_source",
        "citation_mismatch",
        "answer_overreach",
        "insufficient_selected_context",
        "conflicting_contexts",
        "failure_to_abstain",
        "not_observable",
    ],
    "citation_state": [
        "not_applicable",
        "correct",
        "partial",
        "wrong_source",
        "missing",
        "malformed",
    ],
    "source_condition": [
        "current_canonical",
        "current_noncanonical",
        "stale",
        "superseded",
        "low_authority",
        "conflicting_authorities",
        "unknown",
    ],
    "abstention_requirement": [
        "must_answer",
        "may_answer_with_qualification",
        "must_abstain",
    ],
    "evidence_role": ["supporting", "contradicting"],
    "confidence": ["1", "2", "3", "4", "5"],
    "complete": ["yes", "no"],
}

LABEL_HEADERS = [
    "Case ID",
    "Claim #",
    "Claim text",
    "Answer start",
    "Answer end",
    "Propositional?",
    "Claim verdict",
    "Failure label",
    "Primary root cause",
    "Citation state",
    "Source condition",
    "Abstention requirement",
    "Evidence 1 source ID",
    "Evidence 1 chunk ID",
    "Evidence 1 start",
    "Evidence 1 end",
    "Evidence 1 exact text",
    "Evidence 1 role",
    "Evidence 2 source ID",
    "Evidence 2 chunk ID",
    "Evidence 2 start",
    "Evidence 2 end",
    "Evidence 2 exact text",
    "Evidence 2 role",
    "Confidence 1–5",
    "Short rationale",
    "Complete?",
]


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value
    else:
        text = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if len(text) <= MAX_EXCEL_TEXT:
        return text
    return text[: MAX_EXCEL_TEXT - 30] + "\n[truncated in workbook]"


def _header(ws: Any, row: int, values: Sequence[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _style_table(ws: Any, widths: Sequence[int], freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _instructions_sheet(wb: Workbook, annotator_id: str) -> None:
    ws = wb.active
    ws.title = "START HERE"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 105

    lines = [
        ("ContextTrace 60-case annotation", True),
        (f"This workbook is assigned to {annotator_id}.", False),
        ("", False),
        ("What you do", True),
        ("1. Complete the six practice cases first.", False),
        ("2. Discuss only the practice cases with the other annotator.", False),
        ("3. Complete all 60 production cases independently.", False),
        ("4. Complete the ATTESTATION sheet.", False),
        ("5. Return this workbook to SAR. You do not edit any JSON.", False),
        ("", False),
        ("How to work", True),
        (
            "Read a case in PRACTICE CASES or PRODUCTION CASES. Review its "
            "retrieved chunks in the matching EVIDENCE sheet. Enter one "
            "short, independently verifiable claim per row in the matching "
            "LABELS sheet.",
            False,
        ),
        (
            "The workbook provides 20 claim rows per case. Leave unused rows "
            "blank. Contact SAR if one answer needs more than 20 claims.",
            False,
        ),
        (
            "Copy each claim exactly from the answer. You may leave Answer "
            "start and Answer end blank; SAR calculates them after submission.",
            False,
        ),
        (
            "Copy the shortest decisive evidence span exactly and select its "
            "source and chunk IDs. You may leave evidence start and end blank; "
            "SAR calculates them after submission. Use Evidence 2 only when a "
            "second span is necessary.",
            False,
        ),
        ("", False),
        ("Rules", True),
        (
            "Use your own judgment. Do not use an LLM, ContextTrace, another "
            "verifier, web search, system predictions, or Phase 6 metrics.",
            False,
        ),
        (
            "Do not view or discuss the other production submission until SAR "
            "confirms that both original workbooks have been received and "
            "hashed.",
            False,
        ),
        ("Use only the material supplied in this packet.", False),
        (
            "The dropdown values are defined in VALUE GUIDE. Use "
            "not_observable when the trace cannot identify a root cause.",
            False,
        ),
    ]
    for row_index, (text, heading) in enumerate(lines, start=1):
        cell = ws.cell(row=row_index, column=2, value=text)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if heading:
            cell.font = Font(bold=True, size=14, color="244062")
        else:
            cell.font = Font(size=11)
        ws.row_dimensions[row_index].height = 30 if heading else 26


def _value_guide_sheet(wb: Workbook) -> dict[str, str]:
    ws = wb.create_sheet("VALUE GUIDE")
    _header(ws, 1, ["Field", "Allowed value"])
    named_ranges: dict[str, str] = {}
    row = 2
    for field, values in OPTIONS.items():
        start = row
        for value in values:
            ws.cell(row=row, column=1, value=field)
            ws.cell(row=row, column=2, value=value)
            row += 1
        named_ranges[field] = (
            f"{quote_sheetname(ws.title)}!$B${start}:$B${row - 1}"
        )
    _style_table(ws, [30, 48])
    return named_ranges


def _case_sheet(
    wb: Workbook,
    *,
    packet: Path,
    split: str,
    cases: Sequence[dict[str, Any]],
) -> dict[str, int]:
    title = f"{split.upper()} CASES"
    ws = wb.create_sheet(title)
    headers = [
        "Case ID",
        "Track",
        "Domain",
        "Query",
        "Answer",
        "Answer length",
        "Source family",
        "Source IDs",
        "Source conditions",
        "Temporal pair",
        "Citations",
    ]
    _header(ws, 1, headers)
    case_rows: dict[str, int] = {}
    for row_index, case in enumerate(cases, start=2):
        case_id = str(case["case_id"])
        trace = _load_json(packet / str(case["trace_path"]))
        sources = case.get("sources") or []
        conditions = [
            f"{source.get('source_id')}: "
            + ", ".join(source.get("source_conditions") or [])
            for source in sources
        ]
        pair = case.get("source_condition_pair") or {}
        values = [
            case_id,
            case.get("track"),
            case.get("domain_group"),
            trace.get("query"),
            trace.get("answer"),
            len(str(trace.get("answer") or "")),
            case.get("source_family"),
            "\n".join(str(source.get("source_id")) for source in sources),
            "\n".join(conditions),
            _safe_text(pair),
            _safe_text(trace.get("citations") or []),
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=_safe_text(value))
        ws.row_dimensions[row_index].height = 105
        case_rows[case_id] = row_index
    _style_table(ws, [34, 22, 25, 55, 90, 14, 24, 30, 35, 38, 38])
    return case_rows


def _evidence_sheet(
    wb: Workbook,
    *,
    packet: Path,
    split: str,
    cases: Sequence[dict[str, Any]],
) -> None:
    ws = wb.create_sheet(f"{split.upper()} EVIDENCE")
    headers = [
        "Case ID",
        "Rank",
        "Selected?",
        "Chunk ID",
        "Source ID",
        "Chunk text",
    ]
    _header(ws, 1, headers)
    row = 2
    for case in cases:
        trace = _load_json(packet / str(case["trace_path"]))
        selected_ids = set(case.get("selected_context_ids") or [])
        for rank, chunk in enumerate(trace.get("retrieved_chunks") or [], start=1):
            chunk_id = str(chunk.get("id") or "")
            values = [
                case["case_id"],
                rank,
                "yes" if chunk_id in selected_ids else "no",
                chunk_id,
                chunk.get("source_id"),
                chunk.get("text"),
            ]
            for column, value in enumerate(values, start=1):
                ws.cell(row=row, column=column, value=_safe_text(value))
            ws.row_dimensions[row].height = 90
            row += 1
    _style_table(ws, [34, 8, 12, 45, 30, 115])


def _add_dropdown(
    ws: Any,
    *,
    column: int,
    first_row: int,
    last_row: int,
    source: str,
) -> None:
    validation = DataValidation(
        type="list",
        formula1=f"={source}",
        allow_blank=True,
    )
    validation.error = "Choose a value from the dropdown."
    validation.errorTitle = "Invalid value"
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{first_row}:{letter}{last_row}")


def _labels_sheet(
    wb: Workbook,
    *,
    split: str,
    cases: Sequence[dict[str, Any]],
    case_rows: dict[str, int],
    list_ranges: dict[str, str],
) -> None:
    ws = wb.create_sheet(f"{split.upper()} LABELS")
    _header(ws, 1, LABEL_HEADERS)
    row = 2
    case_sheet = f"{split.upper()} CASES"
    for case in cases:
        case_id = str(case["case_id"])
        for claim_number in range(1, CLAIM_SLOTS_PER_CASE + 1):
            id_cell = ws.cell(row=row, column=1, value=case_id)
            id_cell.hyperlink = f"#'{case_sheet}'!A{case_rows[case_id]}"
            id_cell.style = "Hyperlink"
            ws.cell(row=row, column=2, value=claim_number)
            row += 1
    last_row = row - 1

    validation_columns = {
        6: "propositional",
        7: "claim_verdict",
        8: "failure_label",
        9: "primary_root_cause",
        10: "citation_state",
        11: "source_condition",
        12: "abstention_requirement",
        18: "evidence_role",
        24: "evidence_role",
        25: "confidence",
        27: "complete",
    }
    for column, option_name in validation_columns.items():
        _add_dropdown(
            ws,
            column=column,
            first_row=2,
            last_row=last_row,
            source=list_ranges[option_name],
        )

    incomplete_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add(
        f"A2:AA{last_row}",
        FormulaRule(
            formula=['AND($C2<>"",$AA2<>"yes")'],
            fill=incomplete_fill,
        ),
    )
    _style_table(
        ws,
        [
            34,
            9,
            55,
            13,
            13,
            16,
            22,
            26,
            34,
            22,
            28,
            31,
            30,
            40,
            12,
            12,
            55,
            18,
            30,
            40,
            12,
            12,
            55,
            18,
            16,
            55,
            13,
        ],
    )
    ws.freeze_panes = "C2"


def _sources_sheet(
    wb: Workbook, *, packet: Path, cases: Sequence[dict[str, Any]]
) -> None:
    ws = wb.create_sheet("SOURCES")
    headers = [
        "Source ID",
        "Conditions",
        "Published at",
        "Authority basis",
        "Source URL",
        "Local source file",
    ]
    _header(ws, 1, headers)
    sources: dict[str, dict[str, Any]] = {}
    for case in cases:
        for source in case.get("sources") or []:
            sources[str(source["source_id"])] = source
    for row, source_id in enumerate(sorted(sources), start=2):
        source = sources[source_id]
        values = [
            source_id,
            ", ".join(source.get("source_conditions") or []),
            source.get("published_at"),
            source.get("authority_basis"),
            source.get("source_url"),
            source.get("text_path"),
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=_safe_text(value))
        url_cell = ws.cell(row=row, column=5)
        if url_cell.value:
            url_cell.hyperlink = str(url_cell.value)
            url_cell.style = "Hyperlink"
        local_cell = ws.cell(row=row, column=6)
        if local_cell.value:
            local_cell.hyperlink = str(local_cell.value)
            local_cell.style = "Hyperlink"
        ws.row_dimensions[row].height = 45
    _style_table(ws, [30, 28, 24, 75, 55, 45])


def _attestation_sheet(wb: Workbook, annotator_id: str) -> None:
    ws = wb.create_sheet("ATTESTATION")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 85
    ws["A1"] = "Annotation attestation"
    ws["A1"].font = Font(bold=True, size=16, color="244062")
    ws.merge_cells("A1:B1")
    rows = [
        (
            "Statement",
            "You attest that you personally completed these annotations. You "
            "did not use an LLM, ContextTrace, another verifier, web search, "
            "system predictions, Phase 6 metrics, or the other annotator's "
            "production work. You used only the supplied material.",
        ),
        ("Annotator ID", annotator_id),
        ("Type your full name", ""),
        ("Date completed (YYYY-MM-DD)", ""),
        ("I agree (type YES)", ""),
    ]
    for row, (label, value) in enumerate(rows, start=3):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=2).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        ws.row_dimensions[row].height = 65 if row == 3 else 28


def build_workbook(packet: Path, output: Path) -> None:
    assignment = _load_json(packet / "ASSIGNMENT.json")
    annotator_id = str(assignment["annotator_id"])
    cases = assignment["cases"]
    wb = Workbook()
    _instructions_sheet(wb, annotator_id)
    list_ranges = _value_guide_sheet(wb)

    all_cases: list[dict[str, Any]] = []
    for split in ("practice", "production"):
        split_cases = list(cases[split])
        all_cases.extend(split_cases)
        case_rows = _case_sheet(
            wb,
            packet=packet,
            split=split,
            cases=split_cases,
        )
        _evidence_sheet(
            wb,
            packet=packet,
            split=split,
            cases=split_cases,
        )
        _labels_sheet(
            wb,
            split=split,
            cases=split_cases,
            case_rows=case_rows,
            list_ranges=list_ranges,
        )
    _sources_sheet(wb, packet=packet, cases=all_cases)
    _attestation_sheet(wb, annotator_id)
    wb["VALUE GUIDE"].sheet_state = "hidden"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)

    # Reopen once so packaging failures are caught before distribution.
    checked = load_workbook(output, read_only=True, data_only=False)
    expected = {
        "START HERE",
        "PRACTICE CASES",
        "PRACTICE EVIDENCE",
        "PRACTICE LABELS",
        "PRODUCTION CASES",
        "PRODUCTION EVIDENCE",
        "PRODUCTION LABELS",
        "SOURCES",
        "ATTESTATION",
        "VALUE GUIDE",
    }
    if set(checked.sheetnames) != expected:
        raise ValueError("Workbook validation failed: sheet set differs.")
    checked.close()


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--packet", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.output.exists():
        raise SystemExit(f"Output already exists: {args.output}")
    build_workbook(args.packet, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
