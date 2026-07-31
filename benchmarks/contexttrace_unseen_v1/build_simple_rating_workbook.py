"""Build a one-sheet, 20-case evidence-rating workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Mapping, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Install the workbook dependency first: python -m pip install -r "
        "benchmarks/contexttrace_unseen_v1/requirements-annotation-xlsx.txt"
    ) from exc


SELECTION_SEED = "contexttrace-simple-rating-20-v1"
GROUPS = (
    "software_product_documentation",
    "policy_regulatory",
    "support_operational",
    "temporal_source_condition",
)
CASES_PER_GROUP = 5
MAX_CELL_TEXT = 32_000


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _rank(case_id: str) -> str:
    value = f"{SELECTION_SEED}:{case_id}".encode("utf-8")
    return hashlib.sha256(value).hexdigest()


def _group(case: Mapping[str, Any]) -> str:
    if case["track"] == "temporal_source_condition":
        return "temporal_source_condition"
    return str(case["domain_group"])


def select_cases(cases: Sequence[Mapping[str, Any]]) -> list[Mapping[str, Any]]:
    grouped: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for case in cases:
        grouped[_group(case)].append(case)
    if set(grouped) != set(GROUPS):
        raise ValueError("The production assignment does not contain all four groups.")
    selected: list[Mapping[str, Any]] = []
    for group in GROUPS:
        ranked = sorted(grouped[group], key=lambda case: _rank(str(case["case_id"])))
        if len(ranked) < CASES_PER_GROUP:
            raise ValueError(f"Group {group} has fewer than five cases.")
        selected.extend(ranked[:CASES_PER_GROUP])
    return selected


def _safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if len(text) <= MAX_CELL_TEXT:
        return text
    return text[: MAX_CELL_TEXT - 30] + "\n[truncated in workbook]"


def _evidence_text(case: Mapping[str, Any], trace: Mapping[str, Any]) -> str:
    selected_ids = set(case.get("selected_context_ids") or [])
    chunks = list(trace.get("retrieved_chunks") or [])
    selected = [chunk for chunk in chunks if str(chunk.get("id")) in selected_ids]
    if not selected:
        selected = chunks[:5]

    lines: list[str] = []
    for source in case.get("sources") or []:
        conditions = ", ".join(source.get("source_conditions") or []) or "unknown"
        lines.append(
            "SOURCE "
            f"{source.get('source_id')} | condition: {conditions} | "
            f"published: {source.get('published_at') or 'unknown'}"
        )
    pair = case.get("source_condition_pair") or {}
    if pair:
        lines.append(
            "SOURCE PAIR | "
            f"type: {pair.get('pair_type') or 'unknown'} | "
            f"expected condition: {pair.get('expected_source_condition') or 'unknown'}"
        )
    for index, chunk in enumerate(selected, start=1):
        lines.append(
            f"\nEVIDENCE {index} | {chunk.get('id')} | "
            f"source: {chunk.get('source_id')}\n{chunk.get('text') or ''}"
        )
    return _safe_text("\n".join(lines))


def build_workbook(packet: Path, output: Path, reviewer: str, arm: str) -> None:
    assignment = _load_json(packet / "ASSIGNMENT.json")
    cases = select_cases(assignment["cases"]["production"])

    wb = Workbook()
    ws = wb.active
    ws.title = "RATE 20 CASES"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "ContextTrace 20-case evidence rating"
    ws["A1"].font = Font(bold=True, size=16, color="244062")
    ws.merge_cells("A1:H1")
    ws["A2"] = "Assigned to"
    ws["B2"] = reviewer
    ws["C2"] = "Study arm"
    ws["D2"] = arm
    ws["A3"] = "Your task"
    ws["B3"] = (
        "Read the question, answer, and supplied evidence. Choose one rating "
        "from 1 to 5. Add a short reason only if useful."
    )
    ws.merge_cells("B3:H3")
    ws["A4"] = "Rating scale"
    ws["B4"] = (
        "1 = clearly wrong or contradicted; 2 = mostly wrong or materially "
        "unsupported; 3 = cannot determine from the evidence; 4 = mostly "
        "correct with a minor issue; 5 = fully correct and supported."
    )
    ws.merge_cells("B4:H4")
    ws["A5"] = "Rules"
    ws["B5"] = (
        "Use only this sheet. Do not use an LLM, ContextTrace, web search, "
        "system predictions, Phase 6 results, or the other person's ratings."
    )
    ws.merge_cells("B5:H5")
    ws["A6"] = "Name"
    ws["B6"] = ""
    ws["C6"] = "Date (YYYY-MM-DD)"
    ws["D6"] = ""
    ws["E6"] = "Completed independently (type YES)"
    ws["F6"] = ""
    ws.merge_cells("F6:H6")

    headers = [
        "#",
        "Case ID",
        "Group",
        "Question",
        "Answer",
        "Supplied evidence",
        "Rating 1–5",
        "Short reason (optional)",
    ]
    header_row = 8
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for number, case in enumerate(cases, start=1):
        row = header_row + number
        trace = _load_json(packet / str(case["trace_path"]))
        values = [
            number,
            case["case_id"],
            _group(case),
            trace.get("query"),
            trace.get("answer"),
            _evidence_text(case, trace),
            "",
            "",
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=_safe_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 210

    rating_validation = DataValidation(
        type="list",
        formula1='"1,2,3,4,5"',
        allow_blank=False,
    )
    rating_validation.error = "Choose 1, 2, 3, 4, or 5."
    rating_validation.errorTitle = "Invalid rating"
    rating_validation.showErrorMessage = True
    ws.add_data_validation(rating_validation)
    rating_validation.add(f"G{header_row + 1}:G{header_row + len(cases)}")

    widths = [6, 34, 31, 48, 68, 95, 14, 42]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in range(1, 7):
        ws.row_dimensions[row].height = 32
        for cell in ws[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ("A2", "C2", "A3", "A4", "A5", "A6", "C6", "E6"):
        ws[cell].font = Font(bold=True)
    ws.freeze_panes = f"D{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:H{header_row + len(cases)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    wb.save(output)

    checked = load_workbook(output, read_only=False, data_only=False)
    if checked.sheetnames != ["RATE 20 CASES"]:
        raise ValueError("Workbook must contain exactly one worksheet.")
    if checked.active.max_row != 28:
        raise ValueError("Workbook must contain exactly 20 rating rows.")
    checked.close()


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--packet", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--reviewer", required=True)
    parser.add_argument("--arm", required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.output.exists():
        raise SystemExit(f"Output already exists: {args.output}")
    build_workbook(args.packet, args.output, args.reviewer, args.arm)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
